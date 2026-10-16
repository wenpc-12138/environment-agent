# -*- coding: utf-8 -*-
"""M10 报告导出器：docx（python-docx）与 pdf（reportlab）（contracts.md §3.M10）。

- `docx(sections, path)` / `pdf(sections, path)`：渲染并写文件，返回路径。
- `DOCX_EXPORTER` / `PDF_EXPORTER`：等价函数别名（契约测试兼容）。
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from ..models.core import ReportSection


def docx(sections: list[ReportSection], path: str | None = None) -> str:
    """渲染 docx：章节名作 Heading 2，正文作段落；返回文件路径。"""
    from docx import Document

    out = Path(path or 'data/reports/report.docx')
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading('环境监测报告', level=0)
    for sec in sections or []:
        doc.add_heading(sec.name, level=2)
        for para in str(sec.content).split('\n'):
            para = para.strip()
            if not para:
                continue
            if para.startswith('|'):
                doc.add_paragraph(para)
            else:
                doc.add_paragraph(para)
        if sec.refs:
            doc.add_paragraph('引用标准：' + '；'.join(sec.refs))
    doc.save(str(out))
    return str(out)


def pdf(sections: list[ReportSection], path: str | None = None) -> str:
    """渲染 pdf（reportlab platypus）；返回文件路径。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    out = Path(path or 'data/reports/report.pdf')
    out.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    story: list[Any] = [Paragraph('环境监测报告', styles['Title'])]
    for sec in sections or []:
        story.append(Paragraph(sec.name, styles['Heading2']))
        for para in str(sec.content).split('\n'):
            para = para.strip()
            if para:
                story.append(Paragraph(para.replace('|', ' '), styles['BodyText']))
    SimpleDocTemplate(str(out), pagesize=A4).build(story)
    return str(out)


# 契约测试兼容别名
DOCX_EXPORTER = docx
PDF_EXPORTER = pdf
