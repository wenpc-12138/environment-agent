# -*- coding: utf-8 -*-
"""知识库 M2-M4 管线: 文档解析 / 规则自动分类 / 分块 / 轻量向量化

- M2 文档解析: PDF(DOCX/HTML/TXT/MD/XLSX) -> 结构化文本 + 元数据(标准号/机构/日期)
- M2 自动分类: 规则引擎, 输出 domain/item_type/validity/confidence
- M2 分块:    chunk_text() 按段落切块并保留重叠
- M3 向量化:  hash 词袋向量(确定性/离线/numpy), embed_text() + 余弦相似度
"""
import hashlib
import os
import re

# numpy 可选加载; 缺失时回退纯 Python 向量实现(功能等价, 稍慢)
try:
    import numpy as _np
except Exception:
    _np = None

VEC_DIM = 256

# ---------------- 文档解析 ----------------
_EXT_READERS = {
    ".txt": "text", ".md": "text", ".markdown": "text", ".log": "text",
    ".pdf": "pdf", ".docx": "docx", ".doc": "docx",
    ".html": "html", ".htm": "html",
    ".xlsx": "xlsx", ".xls": "xlsx",
}


def extract_text(path):
    """按扩展名解析文档为纯文本, 失败返回空串"""
    ext = os.path.splitext(path)[1].lower()
    kind = _EXT_READERS.get(ext, "text")
    try:
        if kind == "text":
            for enc in ("utf-8", "gbk", "utf-8-sig"):
                try:
                    with open(path, "r", encoding=enc) as f:
                        return f.read()
                except (UnicodeDecodeError, UnicodeError):
                    continue
            return ""
        if kind == "pdf":
            return _read_pdf(path)
        if kind == "docx":
            return _read_docx(path)
        if kind == "html":
            return _read_html(path)
        if kind == "xlsx":
            return _read_xlsx(path)
    except Exception:
        return ""
    return ""


# ---------------- Markdown 结构化转换(入库保存格式) ----------------
def md_from_text(text):
    """把纯文本规范化为 Markdown(段落间空行分隔, 清理 \\r 与多余空行)"""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [l.rstrip() for l in text.split("\n")]
    out, prev_blank = [], True
    for l in lines:
        if not l.strip():
            if not prev_blank:
                out.append("")
            prev_blank = True
        else:
            prev_blank = False
            out.append(l)
    return "\n".join(out).strip()


def _rows_to_md_table(rows):
    """二维列表 -> Markdown 表格(首行表头, 空值补位, 换行转空格)"""
    rows = [[(c or "").strip().replace("\n", " ") for c in r] for r in rows]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    lines = []
    for i, r in enumerate(rows):
        cells = (r + [""] * (w - len(r)))[:w]
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            lines.append("| " + " | ".join("---" for _ in range(w)) + " |")
    return "\n".join(lines)


def extract_markdown(path):
    """按扩展名把文档统一转换为 Markdown 结构化文本(标题/表格/列表),
    作为知识库入库保存的统一格式; 解析失败回退纯文本规范化"""
    ext = os.path.splitext(path)[1].lower()
    kind = _EXT_READERS.get(ext, "text")
    try:
        if kind == "pdf":
            md = _read_pdf_md(path)
        elif kind == "docx":
            md = _read_docx_md(path)
        elif kind == "html":
            md = _read_html_md(path)
        elif kind == "xlsx":
            md = _read_xlsx_md(path)
        else:
            md = _read_text_md(path)
    except Exception:
        md = md_from_text(extract_text(path))
    # 统一清洗 PDF 转 Markdown 的排版问题(页码残留/目录点线/断行合并)
    return clean_md_layout(md or "")


_MD_STRUCT_RE = re.compile(
    r"^(#{1,6}\s|[-*+]\s+|\d+[.、)]\s|\|.*\||>\s|```|[-=]{3,}\s*$|<[a-zA-Z][^>]*>|<!--|!\[|\]\()"
)


def _is_page_no_line(line):
    """页码残留行: - 1 - / ···3 / .......12"""
    s = line.strip()
    if not s:
        return False
    if re.fullmatch(r"-?\s*\d{1,4}\s*-?", s):          # - 1 -
        return True
    if re.fullmatch(r"[.．·]{3,}\s*\d{1,4}", s):        # .....12
        return True
    if re.fullmatch(r"\d{1,4}\s*[.．·]{3,}", s):        # 12.....
        return True
    return False


def _is_toc_line(line):
    """目录点线行: 第一章 总则 ..............1 / 前言 .......... I"""
    s = line.strip()
    if not s or "|" in s:
        return False
    return bool(re.fullmatch(
        r"[\u4e00-\u9fff0-9A-Za-z（）()·、\s]{1,40}[.．·]{4,}\s*[0-9IVXLCivxlc]{1,6}", s))


def _is_md_struct(line):
    """Markdown 结构行(标题/列表/表格/引用/代码块等), 不参与断行合并"""
    s = line.lstrip()
    return bool(s) and bool(_MD_STRUCT_RE.match(s))


def _is_segment_start(s):
    """段首行: 不并入上一段, 开启新段落(第X条 / (一) / 一、)"""
    if re.match(r"^第[一二三四五六七八九十百0-9]+[条章篇]", s):
        return True
    if re.match(r"^[（(][一二三四五六七八九十]+[)）]", s):
        return True
    if re.match(r"^[一二三四五六七八九十]+、", s):
        return True
    return False


def _join_para(a, b):
    """合并 PDF 断行: 前后均为 ASCII 字母数字时补空格, 中文直接拼接"""
    if not a:
        return b
    a_end = a[-1]
    b_start = b[0]
    if a_end.isascii() and a_end.isalnum() and b_start.isascii() and b_start.isalnum():
        return a + " " + b
    return a + b


def _looks_like_pdf_dump(lines):
    """疑似 PDF 转换产物: 页码残留>=2 或目录点线>=2 或连续无空行普通行>=5"""
    page_cnt = sum(1 for l in lines if _is_page_no_line(l))
    toc_cnt = sum(1 for l in lines if _is_toc_line(l))
    consec = maxc = 0
    for l in lines:
        if l.strip() and not _is_md_struct(l) and not _is_page_no_line(l) and not _is_toc_line(l):
            consec += 1
            maxc = max(maxc, consec)
        else:
            consec = 0
    return page_cnt >= 2 or toc_cnt >= 2 or maxc >= 5


def clean_md_layout(md):
    """清洗 PDF 转 Markdown 的排版问题: 页码残留/目录点线/断行合并。
    页码与目录行始终清除; 断行合并仅在疑似 PDF 转换产物时执行, 避免破坏正常手写 MD。"""
    md = (md or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [l.rstrip() for l in md.split("\n")]
    # 1) 清理页码残留行与目录点线行
    cleaned = [l for l in lines
               if not _is_page_no_line(l) and not _is_toc_line(l)]
    # 2) 非疑似 PDF 产物: 仅清理页码/目录, 其余原样保留
    if not _looks_like_pdf_dump(lines):
        return "\n".join(cleaned).strip()
    # 3) 疑似 PDF 产物: 合并断行(保留 Markdown 结构与列表/条文体)
    out = []
    buf = None

    def _table_gap():
        """表格行后补空行: 防止表格与其后正文(如 9.2标准曲线的建立)无空行
        粘连, 被渲染器误并入表格单元格。out 最后一行为表格行时补一个空行。"""
        if out and "|" in out[-1] and out[-1].strip():
            out.append("")

    for l in cleaned:
        s = l.strip()
        if not s:
            if buf is not None:
                _table_gap()
                out.append(buf)
                buf = None
            continue
        if _is_md_struct(l):
            if buf is not None:
                _table_gap()
                out.append(buf)
                buf = None
            if "|" not in l:
                _table_gap()
            out.append(l)
            continue
        if _is_segment_start(s):
            if buf is not None:
                _table_gap()
                out.append(buf)
            _table_gap()
            buf = s
        elif buf is not None:
            # 断行合并仅针对"短行且非句尾"(即 PDF 折行碎片);
            # 长行(完整段落行)保留独立换行, 保持 PDF 版式可读性
            if len(s) <= 40 and not re.search(r"[。；;:：]$", s):
                buf = _join_para(buf, s)
            else:
                _table_gap()
                out.append(buf)
                buf = s
        else:
            _table_gap()
            buf = s
    if buf is not None:
        _table_gap()
        out.append(buf)
    return "\n".join(out).strip()


def _read_text_md(path):
    """md/markdown 保留原文; 其它文本类规范化为 Markdown 段落"""
    if os.path.splitext(path)[1].lower() in (".md", ".markdown"):
        for enc in ("utf-8", "utf-8-sig", "gbk"):
            try:
                with open(path, "r", encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        return ""
    return md_from_text(extract_text(path))


def _read_pdf_md(path):
    """PDF -> Markdown。

    按《路径A_RapidAI全家桶》优先走纯本地 RapidAI 管线(ocr_worker.py /convert_rapid):
      RapidLayout 版面分析 -> RapidOCR 文字 -> RapidTable 表格结构 -> RapidLaTeXOCR 公式
      -> 按版面坐标拼装 Markdown。纯本地免 VLM/API, 是扫描版 PDF 转 MD 存档的主力方案。
    未安装 RapidAI 全家桶或转换失败 -> 回退 VLM 六步管线(vlm_convert.py):
      纯文本页走文字层快速通道, 扫描/混合页多模态 VLM 视觉理解, 再骨架抽取/结构化重排/质量回校;
      无视觉 Provider 时再回退传统 RapidOCR 路线, 功能不降级。
    所有 fitz/rapidocr/rapid_layout/rapid_table 解析均在独立 OCR worker 进程完成:
    pymupdf 为 AGPL-3.0, 采用与 SearXNG 相同的独立进程隔离,
    主进程不 import 这些库规避许可证传染。
    worker 不可用(未安装/启动失败)时回退 pypdf 纯文本层提取。"""
    # 0) 文字层快速通道: PDF 含足够文字层 且 无表格/公式/图片时直接规范化返回(秒级)。
    #    含富内容(表格/公式/图片)时必须跳过快速通道, 走 RapidAI 完整版面识别,
    #    否则表格结构丢失、公式符号(如 Symbol 字体)变乱码。
    #    纯扫描件(文字层过少)同样不受影响, 仍走后续管线。
    try:
        _lt = _read_pdf(path, with_pages=True)
        if _lt and len(_lt.strip()) >= 300:
            _plain = True
            try:
                import ocr_client
                _rich = ocr_client.probe_rich_pdf(path)
                if _rich is not None:
                    _plain = _rich == 0
            except Exception:
                _plain = True  # worker 不可用则保持原快速通道行为
            if _plain:
                return md_from_text(_lt)
    except Exception:
        pass
    # 1) RapidAI 全家桶(路径A: 纯本地版面+文字+表格+公式 -> MD)
    try:
        import ocr_client
        md, meta = ocr_client.convert_pdf_rapid(path)
        if md and len(md) > 50:
            # 完整性校验：RapidAI 输出页数显著少于 PDF 实际页数(<60%)时，
            # 说明大量页面被跳过(单页异常/引擎失败)，回退 pypdf 纯文本层保内容完整
            # （表格结构可能缺失，但优于缺 17 页正文）
            try:
                _out_pages = int((meta or {}).get("pages") or 0)
                _total_pages = int((meta or {}).get("total") or 0)
                if not _total_pages:
                    from pypdf import PdfReader as _PR
                    _total_pages = len(_PR(path).pages)
                if _total_pages > 4 and _out_pages < _total_pages * 0.6:
                    _plain_full = md_from_text(_read_pdf(path))
                    if _plain_full and len(_plain_full) > len(md):
                        return _plain_full
            except Exception:
                pass
            return md
    except Exception:
        pass
    # 2) VLM 六步管线(扫描件 + 多模态 Provider 时启用)
    try:
        import vlm_convert
        if vlm_convert.supports_vlm():
            md = vlm_convert.convert_pdf_to_md(path)
            if md:
                return md
    except Exception:
        pass
    # 3) 传统 RapidOCR 路线(文字层/表格/标题 + OCR 兜底)
    try:
        import ocr_client
        md = ocr_client.extract_pdf_md(path)
        if md:
            return md
    except Exception:
        pass
    # 4) pypdf 纯文本层(worker 不可用兜底)
    return md_from_text(_read_pdf(path))


def extract_pdf_ocr(path, dpi=200, max_pages=300):
    """扫描版 PDF 全文 OCR: 委托独立 OCR worker 进程(rapidocr_onnxruntime),
    主进程不 import RapidOCR, 失败返回空串。"""
    try:
        import ocr_client
        return ocr_client.ocr_pdf(path, dpi=dpi, max_pages=max_pages)
    except Exception:
        return ""


def _read_docx_md(path):
    """python-docx: 按段落样式映射标题层级, 表格转 Markdown 表格, 保持文档顺序"""
    try:
        import docx
        from docx.oxml.ns import qn
    except Exception:
        return md_from_text(_read_docx(path))
    d = docx.Document(path)
    para_map = {p._p: p for p in d.paragraphs}
    tbl_map = {t._tbl: t for t in d.tables}
    out = []
    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = para_map.get(child)
            if p is None:
                continue
            txt = p.text.strip()
            if not txt:
                continue
            style = (p.style.name or "").lower() if p.style else ""
            if style.startswith("heading"):
                try:
                    lvl = int(style.replace("heading", "").strip() or "1")
                except Exception:
                    lvl = 1
                out.append("#" * min(max(lvl, 1), 6) + " " + txt)
            elif style in ("title", "标题"):
                out.append("# " + txt)
            else:
                out.append(txt)
        elif child.tag == qn("w:tbl"):
            t = tbl_map.get(child)
            if t is None:
                continue
            md = _rows_to_md_table([[c.text.strip() for c in r.cells] for r in t.rows])
            if md:
                out.append(md)
    return "\n\n".join(out).strip()


def _soup_to_md(soup):
    """BeautifulSoup 对象 -> Markdown(h1-h6/p/列表/表格/链接/图片), 网页与 HTML 文件共用"""
    for tag in soup(["script", "style", "nav", "footer", "form", "noscript", "iframe"]):
        tag.decompose()
    out = []

    def _inline(el):
        return el.get_text(" ", strip=True) if el else ""

    def _walk(node):
        for el in getattr(node, "children", []):
            name = getattr(el, "name", None)
            if name is None:
                t = (getattr(el, "string", "") or "").strip()
                if t:
                    out.append(t)
                continue
            if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
                t = _inline(el)
                if t:
                    out.append("#" * int(name[1]) + " " + t)
            elif name == "p":
                t = _inline(el)
                if t:
                    out.append(t)
            elif name in ("ul", "ol"):
                for i, li in enumerate(el.find_all("li", recursive=False)):
                    t = _inline(li)
                    if t:
                        out.append(("- " if name == "ul" else "%d. " % (i + 1)) + t)
            elif name == "table":
                rows = []
                for tr in el.find_all("tr"):
                    cells = [_inline(td) for td in tr.find_all(["td", "th"])]
                    if cells:
                        rows.append(cells)
                md = _rows_to_md_table(rows)
                if md:
                    out.append(md)
            elif name == "a":
                t = _inline(el)
                href = (el.get("href") or "").strip()
                if t:
                    out.append("[%s](%s)" % (t, href) if href else t)
            elif name == "img":
                src = (el.get("src") or "").strip()
                # 过滤站内装饰/背景图(/images/ 前缀)与内联 base64 图
                if src and not src.lower().startswith(("/images/", "data:image", "data:,")):
                    out.append("![%s](%s)" % ((el.get("alt") or "").strip(), src))
            elif name == "hr":
                out.append("---")
            elif name in ("div", "section", "article", "main", "body", "blockquote",
                          "li", "dl", "dd", "dt", "thead", "tbody", "tr", "br"):
                _walk(el)
    _walk(soup.body or soup)
    res, prev_blank = [], True
    for l in out:
        if not l.strip():
            if not prev_blank:
                res.append("")
            prev_blank = True
        else:
            prev_blank = False
            res.append(l)
    return "\n".join(res).strip()


def _read_html_md(path):
    """HTML 文件 -> Markdown(结构化), 复用 _soup_to_md"""
    from bs4 import BeautifulSoup
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f.read(), "lxml")
    return _soup_to_md(soup)


# 正规来源域名后缀白名单: 中国政府机构(gov.cn)及港澳台官方域名
_TRUSTED_URL_SUFFIXES = (".gov.cn", ".gov.hk", ".gov.mo", ".gov.tw")


def is_trusted_url(url):
    """正规来源校验: 仅放行官方机构域名, 且解析地址非内网/回环/保留段(防 SSRF)"""
    from urllib.parse import urlparse
    host = (urlparse(url).hostname or "").lower()
    if not host:
        return False
    if not any(host == s.lstrip(".") or host.endswith(s) for s in _TRUSTED_URL_SUFFIXES):
        return False
    try:
        import ipaddress
        import socket
        ip = ipaddress.ip_address(socket.gethostbyname(host))
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
            return False
    except Exception:
        return False
    return True


def url_to_markdown(url, timeout=15):
    """抓取网页正文转 Markdown: 正规来源校验 -> requests 抓取 -> _soup_to_md -> 排版清洗
    返回 (markdown, title_hint); 来源不正规或抓取/解析失败抛异常, 由调用方记录处理"""
    if not is_trusted_url(url):
        raise ValueError("来源非正规官方域名(仅允许 gov.cn 等), 已拒绝导入: %s" % url)
    import requests
    hdr = {
        "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
                       " (KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"),
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
    }
    resp = requests.get(url, headers=hdr, timeout=timeout, verify=True)
    resp.raise_for_status()
    # 编码兜底: 无 charset 声明时 requests 默认 ISO-8859-1 会乱码,
    # 用 chardet(apparent_encoding) 检测真实编码后重新解码
    enc = resp.encoding
    if enc is None or enc.lower() in ("iso-8859-1", "latin-1", "ascii"):
        ae = resp.apparent_encoding
        if ae and ae.lower() not in ("iso-8859-1", "latin-1", "ascii"):
            enc = ae
    html = resp.content.decode(enc or "utf-8", errors="replace")
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "lxml")
    # 标题提示: 优先正文 h1, 回退 <title>
    title_hint = ""
    for h1 in soup.find_all("h1"):
        t = h1.get_text(" ", strip=True)
        if t:
            title_hint = t
            break
    if not title_hint and soup.title and soup.title.get_text(strip=True):
        title_hint = soup.title.get_text(strip=True)
    md = _soup_to_md(soup)
    return clean_md_layout(md), title_hint


def _read_xlsx_md(path):
    """openpyxl: 每个工作表 -> Markdown 表格(首行表头)"""
    try:
        from openpyxl import load_workbook
    except Exception:
        return md_from_text(_read_xlsx(path))
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            while cells and cells[-1] == "":
                cells.pop()
            if cells:
                rows.append(cells)
        if rows:
            md = _rows_to_md_table(rows)
            if md:
                parts.append("## " + (ws.title or "Sheet") + "\n\n" + md)
    wb.close()
    return "\n\n".join(parts).strip()


def save_md_copy(md_dir, eid, md_text):
    """把 Markdown 正文落盘到 kb/ 目录(数据目录), 返回文件路径(失败返回空串)"""
    try:
        os.makedirs(md_dir, exist_ok=True)
        p = os.path.join(md_dir, "%04d.md" % eid)
        with open(p, "w", encoding="utf-8") as f:
            f.write(md_text or "")
        return p
    except Exception:
        return ""


def _read_pdf(path, with_pages=False):
    # 纯文本层提取用 pypdf (BSD); 表格/标题还原与 OCR 走独立 worker 见 _read_pdf_md
    # with_pages=True: 逐页提取并在页边界插入分页标记(--- + <!-- PDF 第 N 页 -->)
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        if not with_pages:
            return "\n".join(p.extract_text() or "" for p in reader.pages)
        parts = []
        for i, p in enumerate(reader.pages, 1):
            t = (p.extract_text() or "").strip()
            if not t:
                continue
            if i > 1:
                parts.append("---\n\n<!-- PDF 第 %d 页 -->" % i)
            parts.append(t)
        return "\n\n".join(parts)
    except Exception:
        return ""


def _read_docx(path):
    try:
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for tbl in d.tables:
            for row in tbl.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _read_html(path):
    try:
        from bs4 import BeautifulSoup
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            soup = BeautifulSoup(f.read(), "lxml")
        for tag in soup(["script", "style", "nav", "footer"]):
            tag.decompose()
        text = soup.get_text("\n")
        return re.sub(r"\n{3,}", "\n\n", text).strip()
    except Exception:
        return ""


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        wb.close()
        return "\n".join(parts)
    except Exception:
        return ""


# ---------------- 元数据规则提取 ----------------
_STD_RE = re.compile(
    r"([A-Z]{1,4}(?:[/-][A-Z]{1,4})?/?(?:T|Z)?[-—–]?\s*\d{2,6}(?:\.\d+)*(?:[-—–]\d{4})?)")
_DATE_RE = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日")
_DATE_RE2 = re.compile(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})")

_ORG_RULES = [
    (r"生态环境部|环境保护部|国家环保局|环保部", "生态环境部"),
    (r"住房和城乡建设部|住建部|建设部", "住房和城乡建设部"),
    (r"市场监督管理总局|质检总局", "国家市场监督管理总局"),
    (r"工业和信息化部|工信部", "工业和信息化部"),
    (r"交通运输部", "交通运输部"),
    (r"卫生健康委|卫生计生委|卫生部", "国家卫生健康委员会"),
    (r"应急管理部", "应急管理部"),
    (r"自然资源部|国土资源部", "自然资源部"),
    (r"水利部", "水利部"),
    (r"农业农村部|农业部", "农业农村部"),
    (r"能源局|发改委|国家发展改革委", "国家发展改革委"),
    (r"国家电网|电力", "国家电网"),
]

_STD_DOMAIN = {
    "GB": "国家标准", "GBJ": "国家标准", "GBZ": "国家标准",
    "HJ": "环保行业", "HJ/T": "环保行业",
    "ISO": "国际标准", "IEC": "国际标准",
    "DL": "电力行业", "DL/T": "电力行业",
    "CJ": "城镇建设", "CJJ": "城镇建设",
    "DB": "地方标准", "DBJ": "地方标准",
    "TB": "铁路行业", "JGJ": "建筑工程", "JG/T": "建筑工程",
    "TSG": "特种设备", "WS": "卫生健康", "SL": "水利行业",
}

_STD_ITEM_TYPE = {
    "standard": ("GB", "HJ", "ISO", "DL", "CJ", "DB", "TB", "JG", "SL", "WS", "TSG", "IEC"),
    "regulation": ("办法", "条例", "规定", "实施细则", "指导意见", "规划", "纲要"),
    "policy": ("通知", "公告", "函", "复函", "批复"),
    "manual": ("手册", "指南", "导则", "技术规范", "操作规程", "作业指导"),
}

_DOMAIN_RULES = [
    (r"环境影响评价|环评|验收监测|竣工验收", "环评"),
    (r"环境保护法|水污染防治|大气污染防治|土壤污染防治|噪声|固废|排污许可", "法律"),
    (r"排放标准|限值|监测技术|分析方法|采样", "工程"),
    (r"应急预案|风险管理|隐患排查|应急演练", "应急"),
    (r"碳达峰|碳中和|温室气体|碳交易", "双碳"),
]

_VALIDITY_RE = re.compile(r"(现行|现行有效|被替代|废止|代替.*?号|已废止)")


def parse_meta(title, text, url=""):
    """规则提取结构化元数据: std_no/org/publish_date/implement_date/validity"""
    blob = (title or "") + "\n" + (text or "")[:2000]
    meta = {"std_no": "", "org": "", "publish_date": "", "implement_date": "", "validity": ""}
    m = _STD_RE.search(blob)
    if m:
        meta["std_no"] = m.group(1).replace(" ", "")
    for pat, org in _ORG_RULES:
        if re.search(pat, blob):
            meta["org"] = org
            break
    m = _DATE_RE.search(blob)
    if m:
        meta["publish_date"] = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    else:
        m = _DATE_RE2.search(blob)
        if m:
            meta["publish_date"] = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = _VALIDITY_RE.search(blob)
    if m:
        v = m.group(0)
        if "废止" in v or "被替代" in v:
            meta["validity"] = "被替代" if "被替代" in v else "废止"
        else:
            meta["validity"] = "现行"
    return meta


def classify_doc(title, text="", url=""):
    """规则自动分类: 返回 (domain, item_type, source_type, confidence)"""
    blob = (title or "") + " " + (text or "")[:800]
    std_no = ""
    m = _STD_RE.search(blob)
    if m:
        std_no = m.group(1).replace(" ", "")
    # 0. 征求意见稿优先归类: 任何出现"征求意见稿/征求意见"的文档统一归为征求意见稿
    if "征求意见稿" in blob or "（征求意见稿）" in blob or re.search(r"征求.{0,6}意见", blob):
        domain = "综合"
        for pat, d in _DOMAIN_RULES:
            if re.search(pat, blob):
                domain = d
                break
        return domain, "draft", "政府文件" if "gov.cn" in (url or "") else "原文", 0.95
    # 1. 标准号判定 -> 标准类 + 标准领域
    if std_no:
        for prefix, domain in _STD_DOMAIN.items():
            if std_no.startswith(prefix):
                return domain, "standard", "原文", 0.95
        return "标准", "standard", "原文", 0.9
    # 2. 法规/政策/手册关键词 -> item_type
    itype = "knowledge"
    for ws, t in (("办法 条例 规定 实施细则 指导意见 规划", "regulation"),
                  ("通知 公告 意见 方案", "policy"),
                  ("指南 导则 技术规范 手册 操作规程 作业指导书", "manual")):
        if any(w in blob for w in ws.split()):
            itype = t
            break
    # 3. 领域关键词
    domain = "综合"
    for pat, d in _DOMAIN_RULES:
        if re.search(pat, blob):
            domain = d
            break
    # 4. 来源类型
    if "gov.cn" in (url or ""):
        source_type = "政府文件"
    elif url:
        source_type = "网络"
    else:
        source_type = "原文"
    return domain, itype, source_type, 0.7


# ---------------- LLM 兜底分类 ----------------
def classify_llm(title, text="", fallback=None):
    """LLM 兜底分类: 规则置信度不足时调用活跃 Provider(OpenAI 兼容)补齐分类;
    调用失败或未配置真实 Provider 时回退规则结果。
    返回 (domain, item_type, source_type, confidence)"""
    if fallback is None:
        fallback = classify_doc(title, text)
    fd, fi, fs, fc = fallback
    try:
        import db
        import providers
        prov = db.get_active_provider()
        if not prov or prov.get("kind") == "mock":
            return fallback  # 未配置真实 Provider, 直接回退规则结果
        prompt = (
            "你是生态环境标准/资料分类助手。根据下列文档标题与摘要,输出 JSON:"
            '{"domain":"领域","item_type":"类型","source_type":"来源"}。\n'
            "domain 取值: 环评/法律/工程/应急/双碳/标准/综合;"
            "item_type 取值: standard/regulation/policy/manual/knowledge/draft(征求意见稿);"
            "source_type 取值: 原文/政府文件/网络。\n"
            f"标题: {title}\n摘要: {(text or '')[:1200]}\n"
            "仅输出 JSON,不要其它解释。"
        )
        msgs = [{"role": "user", "content": prompt}]
        acc = []
        p = providers.build_provider(prov)
        for chunk in p.stream(msgs, tools=None, temperature=0.1):
            d = chunk.get("delta") or {}
            if d.get("content"):
                acc.append(d["content"])
        out = "".join(acc).strip()
        import json as _json
        m = re.search(r"\{.*\}", out, re.S)
        if not m:
            return fallback
        obj = _json.loads(m.group(0))
        domain = str(obj.get("domain") or fd)
        itype = str(obj.get("item_type") or fi)
        if itype not in ("standard", "regulation", "policy", "manual", "knowledge", "draft"):
            itype = fi
        stype = str(obj.get("source_type") or fs)
        # LLM 兜底置信度 = 规则置信度提升到 0.9(仅当输出结构与预期一致)
        return domain, itype, stype, 0.9
    except Exception:
        return fallback


# ---------------- 分块 ----------------
def chunk_text(text, max_len=600, overlap=80):
    """按段落切块, 超长段落按句子边界二次切分, 返回 [(chunk, type)]"""
    if not text:
        return []
    text = re.sub(r"\r\n", "\n", text).strip()
    if not text:
        return []
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(text) <= max_len:
        return [(text, "para")]
    chunks = []
    buf = ""
    for p in paras:
        if len(p) > max_len:
            if buf:
                chunks.append((buf, "para"))
                buf = ""
            chunks.extend(_split_long(p, max_len, overlap))
            continue
        if len(buf) + len(p) + 1 <= max_len:
            buf = (buf + "\n" + p) if buf else p
        else:
            chunks.append((buf, "para"))
            buf = p
    if buf:
        chunks.append((buf, "para"))
    return chunks


def _split_long(para, max_len=600, overlap=80):
    """超长段落按句子边界切分, 保留 overlap 重叠"""
    sents = re.split(r"(?<=[。；;!?])", para)
    out, buf = [], ""
    for s in sents:
        s = s.strip()
        if not s:
            continue
        if len(buf) + len(s) <= max_len:
            buf += s
        else:
            if buf:
                out.append((buf, "sentence"))
            tail = buf[-overlap:] if len(buf) > overlap else buf
            buf = tail + s
    if buf:
        out.append((buf, "sentence"))
    return out


# ---------------- 向量化 ----------------
# 语义嵌入模型单例(懒加载; bge-small-zh-v1.5 输出 512 维, 与 VEC_DIM 不同,
# 使用语义模型时统一升维/截断到 VEC_DIM 以兼容 vec0 建表)
_SEM_MODEL = None
_SEM_EMBEDDING_DIM = None
SEM_MODEL_NAME = "BAAI/bge-small-zh-v1.5"
SEM_MODEL_DIM = 512


def _ensure_semantic_model():
    """懒加载 fastembed 语义模型; 失败返回 None(降级哈希词袋)"""
    global _SEM_MODEL, _SEM_EMBEDDING_DIM
    if _SEM_MODEL is not None:
        return _SEM_MODEL
    try:
        from fastembed import TextEmbedding
        _SEM_MODEL = TextEmbedding(SEM_MODEL_NAME, cache_dir=_sem_cache_dir())
        _SEM_EMBEDDING_DIM = SEM_MODEL_DIM
        return _SEM_MODEL
    except Exception:
        _SEM_MODEL = False
        return None


def _sem_cache_dir():
    """模型缓存目录(数据目录下, 便于打包/迁移)"""
    try:
        import db
        return os.path.join(db.DATA_DIR, "kb", "models")
    except Exception:
        return None


def embed_text_semantic(text):
    """语义嵌入(优先 bge-small-zh, 失败回退哈希词袋); 输出维度固定 VEC_DIM"""
    model = _ensure_semantic_model()
    if model is not False and model is not None:
        try:
            vec = next(model.embed([text or " "]))
            arr = _np.asarray(vec, dtype="float32")
            # 维度对齐: >VEC_DIM 截断, <VEC_DIM 补零
            if arr.size != VEC_DIM:
                tmp = _np.zeros(VEC_DIM, dtype="float32")
                n = min(arr.size, VEC_DIM)
                tmp[:n] = arr[:n]
                arr = tmp
            nrm = float(_np.linalg.norm(arr)) or 1.0
            return arr / nrm
        except Exception:
            pass
    return embed_text(text)


def _tokens(text):
    """轻量分词: 中文 bigram + 英文单词, 全部小写"""
    text = (text or "").lower()
    toks = []
    # 中文连续片段 -> bigram
    for seg in re.findall(r"[\u4e00-\u9fff]+", text):
        if len(seg) == 1:
            toks.append(seg)
        else:
            toks.extend(seg[i:i + 2] for i in range(len(seg) - 1))
            toks.extend((seg[i], seg[-2:]) for i in [0]) if False else None
    # 英文/数字词
    toks.extend(re.findall(r"[a-z0-9]{2,}", text))
    return toks


def embed_text(text, dim=VEC_DIM):
    """哈希词袋向量(确定性, 离线), 返回归一化列表或 numpy 数组"""
    vec = [0.0] * dim
    for t in _tokens(text):
        h = int(hashlib.md5(t.encode("utf-8")).hexdigest()[:8], 16)
        vec[h % dim] += 1.0
    norm = (sum(v * v for v in vec) ** 0.5) or 1.0
    vec = [v / norm for v in vec]
    if _np is not None:
        return _np.array(vec, dtype="float32")
    return vec


def cosine(a, b):
    """余弦相似度(兼容 list/np.ndarray)"""
    if _np is not None:
        return float(_np.dot(a, b))
    return sum(x * y for x, y in zip(a, b))
